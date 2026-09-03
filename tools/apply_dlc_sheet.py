# -*- coding: utf-8 -*-
"""Đưa bản dịch DLC (AOC) từ workbook riêng vào bundle của title AOC.

## DLC nằm ở một title khác, patch cũ không chạm tới

Nội dung DLC 1 (`UNLOGICAL [010068501FF9B001][v0][DLC 1].nsp`, AOC 2,25 MB) không nằm
trong `scenario01` của game gốc mà trong romfs của chính AOC — bóc bằng
`extract_dlc_romfs.py` (mượn crypto của `extract_exefs.py`):

    scenario/scenario_aoc01   ScenarioData riêng: 5 scenario 1005..1009 (script 09_01..09_05,
                              "朝のひと時" — mỗi nhân vật một truyện ngắn), 399 câu, 0 lựa chọn,
                              0 tag hiển thị; scenariolist, ChapterAlready
    json/json_aoc01           DLCData_01: 5 charaName (chữ hiện ở danh sách nhân vật của màn
                              Download Contents) + tên sprite
    sprite/sprite_jp_aoc01    16 ảnh; texture/texture_aoc01: nền

LayeredFS áp theo title, nên mod cho `010068501ff9a000` không đổi được gì ở đây: cần một
romfs mod thứ hai đặt dưới `contents/010068501ff9b001/`. Tool này ghi ra bản làm việc
`WORK_AOC` (junction vào Ryujinx như bản gốc). Sheet gốc `UNLOGICAL_v2` không có tab DLC;
bản dịch DLC nằm trong workbook riêng `UNLOGICAL_DLC1.xlsx`, tab `sd_1005..sd_1009`, cùng
khuôn 4 cột ID / Speaker / Japanese / Vietnamese, id `1005/txt/0000`.

## Cách ghi

- Luôn xuất phát từ bundle GỐC của AOC (`STOCK_AOC`), nên chạy lại bao nhiêu lần cũng
  cho cùng kết quả — không có ba chiều, không có "build đã đổi".
- Chốt từng ô như `apply_sheet_cells.py`: JP của sheet phải khớp `text[j]` gốc (làm
  phẳng khoảng trắng), multiset tag `[...]` khớp (trừ `[主人公]`), `[主人公]` cùng có hoặc
  cùng không — trừ ô trong `TOKEN_DROP_OK` có ghi lý do, và chỉ khi ô đó không phải bản
  đôi tên mặc định (`isDefaultNameAdjust`/`isCustomNameAdjust` đều False); ô
  `isDefaultNameAdjust` phải mang tên mặc định `Kanna` bằng chữ.
- `scriptText` (bản sao không ai index) soi theo cùng cách `mirror()`.
- Nameplate `talkName` không có trên sheet (quy ước chung, xem CLAUDE.md): đổi theo đúng
  dạng bản gốc đang dùng `【khoá Nhật/tên La-tinh】` — `NAMEPLATE`, đếm từ build gốc
  (`【雅火/Miyabi】` 1936 lần, …). Khoá Nhật trước dấu `/` là khoá tra voice/sprite, giữ.
- `DLCData_01.charaName.jp`: 5 tên hiện ở danh sách — đổi sang tên La-tinh; các field
  sprite bên cạnh giữ nguyên vì là khoá.

    python tools\\apply_dlc_sheet.py            # chạy thử
    python tools\\apply_dlc_sheet.py --apply    # ghi WORK_AOC
    python tools\\apply_dlc_sheet.py --check    # WORK_AOC không còn kana, nameplate đúng
"""
import io
import json
import os
import re
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

import UnityPy   # noqa: E402
from openpyxl import load_workbook   # noqa: E402

# `--dlc N` (mặc định 1). DLC 2 = title 010068501FF9B002, cùng khung: scenario 2005..2009,
# script 09_06..09_10 ("デート"), thêm mỗi scenario MỘT lựa chọn ba phương án trong `selText`
# — sheet DLC ghi ba phương án đó ở hàng `{sid}/cmd/0000..0002`, đúng thứ tự trong JSON.
DLC = int(sys.argv[sys.argv.index("--dlc") + 1]) if "--dlc" in sys.argv else 1
TITLE_ID = "010068501ff9b00%d" % DLC
STOCK_AOC = r"D:\Downloads\UNLOGICAL_DLC%d\romfs" % DLC     # dump gốc của AOC (extract_dlc_romfs.py)
WORK_AOC = r"D:\Downloads\%s\romfs" % TITLE_ID              # bản làm việc = mod cho title AOC
SHEET = r"D:\Downloads\UNLOGICAL_DLC%d.xlsx" % DLC
SCEN = os.path.join("scenario", "scenario_aoc0%d" % DLC)
JSONB = os.path.join("json", "json_aoc0%d" % DLC)
DLCDATA = "DLCData_0%d" % DLC
APPLY = "--apply" in sys.argv
CHECK = "--check" in sys.argv

# dạng nameplate bản gốc đang dùng (đếm trên ScenarioData của build, 02/09/2026)
NAMEPLATE = {
    "【雅火】": "【雅火/Miyabi】",
    "【宗像 戒】": "【宗像 戒/Munakata Kai】",
    "【永守 藍】": "【永守 藍/Nagamori Ran】",
    "【弥坂 奏壱】": "【弥坂 奏壱/Yasaka Soichi】",
    "【神楽 侑莉/ユーリ】": "【神楽 侑莉/Yuri】",
    "【player】": "【player】",
}
CHARA_NAME = {"雅火": "Miyabi", "宗像 戒": "Munakata Kai", "永守 藍": "Nagamori Ran",
              "弥坂 奏壱": "Yasaka Soichi", "ユーリ": "Yuri"}
# Ô bản Nhật có [主人公] mà bản dịch bỏ tên, thay bằng đại từ. Chỉ cho qua khi ô không
# phải bản đôi tên mặc định (cả hai cờ False), vì bản đôi mà lệch là in sai một trong hai.
TOKEN_DROP_OK = {
    # `[主人公]がそっちに向かおうとする` -> "em định đi ra chỗ đó": kể chuyện ngôi thứ nhất
    # của Kai, tên nhân vật thành "em" là đúng giọng; ô này không có bản đôi.
    "1006/txt/0004",
    # DLC 2 — `恋人になった[主人公]に対して…` -> "Khi đã thành người yêu của nhau, tôi cũng…":
    # lời kể của Ran, bản dịch chuyển sang "của nhau"; ô không có bản đôi.
    "2007/txt/0100",
}
DEFAULT_GIVEN = "Kanna"

TAG = re.compile(r"\[[^\[\]\n]*\]")
KANA = re.compile(r"[ぁ-ゖァ-ヺ]")
SD_ID = re.compile(r"^(\d+)/txt/(\d+)$")
CMD_ID = re.compile(r"^(\d+)/cmd/(\d+)$")     # trên sheet DLC: phương án thứ k của lựa chọn


def flat(s):
    return re.sub(r"[\s\u3000]+", " ", s or "").strip()


def load_text_asset(path, name):
    env = UnityPy.load(path)
    for o in env.objects:
        if o.type.name == "TextAsset" and o.read().m_Name == name:
            d = o.read()
            raw = d.m_Script
            if not isinstance(raw, str):
                raw = bytes(raw).decode("utf-8")
            return env, d, raw
    raise SystemExit("không thấy %s trong %s" % (name, path))


def mirror(script, old, new):
    lines, ol, nl = script.split("\n"), old.split("\n"), new.split("\n")
    hits = [k for k in range(len(lines) - len(ol) + 1) if lines[k:k + len(ol)] == ol]
    if len(hits) != 1:
        return None
    k = hits[0]
    lines[k:k + len(ol)] = nl
    return "\n".join(lines)


def read_sheet():
    wb = load_workbook(SHEET, read_only=True, data_only=True)
    cells = {}
    for ws in wb.worksheets:
        if not ws.title.startswith("sd_"):
            continue
        for row in ws.iter_rows(values_only=True):
            if not row or not isinstance(row[0], str):
                continue
            if not (SD_ID.match(row[0].strip()) or CMD_ID.match(row[0].strip())):
                continue
            jp = row[2] if len(row) > 2 and isinstance(row[2], str) else ""
            vi = row[3] if len(row) > 3 and isinstance(row[3], str) else ""
            cells[row[0].strip()] = (jp, vi.replace("\r\n", "\n"))
    wb.close()
    return cells


def check_work():
    env, d, raw = load_text_asset(os.path.join(WORK_AOC, SCEN), "ScenarioData")
    data = json.loads(raw.lstrip("\ufeff"))
    bad = []
    for t in data["target"]:
        for j, s in enumerate(t["text"]):
            if isinstance(s, str) and KANA.search(TAG.sub("", s)):
                bad.append("%s/txt/%04d còn kana: %r" % (t["scenarioID"], j, s[:40]))
        for k, s in enumerate(t["selText"]):
            if s and KANA.search(s):
                bad.append("%s/selText[%d] còn kana: %r" % (t["scenarioID"], k, s[:60]))
        for n in t["talkName"]:
            if n and n not in NAMEPLATE.values():
                bad.append("nameplate lạ %r" % n)
    _, _, rawj = load_text_asset(os.path.join(WORK_AOC, JSONB), DLCDATA)
    for k in CHARA_NAME:
        if '"jp": "%s"' % k in rawj:
            bad.append("%s còn charaName %r" % (DLCDATA, k))
    for b in bad[:10]:
        print("  FAIL", b)
    if bad:
        raise SystemExit("%d lỗi trong %s" % (len(bad), WORK_AOC))
    print("PASS %s: 0 ô kana (text + selText), nameplate đúng dạng, %s đã La-tinh" % (WORK_AOC, DLCDATA))


def main():
    if CHECK:
        check_work()
        return
    cells = read_sheet()
    env, d, raw = load_text_asset(os.path.join(STOCK_AOC, SCEN), "ScenarioData")
    bom = "\ufeff" if raw.startswith("\ufeff") else ""
    data = json.loads(raw.lstrip("\ufeff"))
    print("sheet: %d ô | AOC: %d scenario, %d câu" % (
        len(cells), len(data["target"]), sum(len(t["text"]) for t in data["target"])))

    def replace_array(blob, old_list, new_list, what):
        """Thay nguyên một mảng JSON; thử cả hai kiểu phân cách, đòi khớp đúng một lần.

        Không thay theo từng chuỗi: câu một dòng xuất hiện y hệt ở `text[]` lẫn
        `scriptText_Line[]` (bản thô của script), thay mù là đụng cả bản thô."""
        for sep in ((",", ":"), (", ", ": ")):
            oj = json.dumps(old_list, ensure_ascii=False, separators=sep)
            if blob.count(oj) == 1:
                return blob.replace(oj, json.dumps(new_list, ensure_ascii=False, separators=sep))
        raise SystemExit("%s không khớp đúng một lần trong file" % what)

    out = raw
    n_ok = n_name = n_sel = 0
    problems = []
    mirrored = mirror_failed = 0
    for ti, t in enumerate(data["target"]):
        sid = t["scenarioID"]
        script = cur_s = t["scriptText"]
        new_text = list(t["text"])
        for j, cur in enumerate(t["text"]):
            key = "%d/txt/%04d" % (sid, j)
            if key not in cells:
                if cur.strip():
                    problems.append((key, "không có trên sheet"))
                continue
            jp, vi = cells[key]
            if flat(jp) != flat(cur):
                problems.append((key, "JP sheet khác build: %r != %r" % (flat(jp)[:30], flat(cur)[:30])))
                continue
            if not vi.strip():
                problems.append((key, "VN trống"))
                continue
            tj = sorted(x for x in TAG.findall(cur) if x != "[主人公]")
            tv = sorted(x for x in TAG.findall(vi) if x != "[主人公]")
            if tj != tv:
                problems.append((key, "tag lệch %s != %s" % (tj, tv)))
                continue
            has_tok = "[主人公]" in cur, "[主人公]" in vi
            paired = t["isDefaultNameAdjust"][j] or t["isCustomNameAdjust"][j]
            if has_tok[0] != has_tok[1]:
                if key in TOKEN_DROP_OK and not paired:
                    print("~  %s: bỏ [主人公] theo TOKEN_DROP_OK" % key)
                else:
                    problems.append((key, "[主人公] chỉ có một bên"))
                    continue
            if t["isDefaultNameAdjust"][j] and DEFAULT_GIVEN not in vi:
                problems.append((key, "ô tên mặc định không có %r" % DEFAULT_GIVEN))
                continue
            if KANA.search(TAG.sub("", vi)):
                problems.append((key, "VN còn kana"))
                continue
            new_text[j] = vi
            nxt = mirror(cur_s, cur, vi)
            if nxt is None:
                mirror_failed += 1
            else:
                cur_s = nxt
                mirrored += 1
            n_ok += 1
        if new_text != t["text"]:
            out = replace_array(out, t["text"], new_text, "text[] target[%d]" % ti)
        # lựa chọn: selText[k] là JSON lồng `{"target":[…]}`; sheet ghi phương án i ở
        # `{sid}/cmd/{i:04d}`. Thay cả chuỗi selText[k] (như apply_sheet_cells), không thay
        # từng phương án — hai phương án có thể trùng chữ.
        for k, cur_raw in enumerate(t["selText"]):
            if not cur_raw:
                continue
            doc = json.loads(cur_raw)
            opts = doc["target"]
            new_opts = list(opts)
            for i, opt in enumerate(opts):
                key = "%d/cmd/%04d" % (sid, i)
                if key not in cells:
                    problems.append((key, "phương án không có trên sheet: %r" % opt[:30]))
                    continue
                jp, vi = cells[key]
                if flat(jp) != flat(opt):
                    problems.append((key, "JP sheet khác build: %r != %r" % (flat(jp)[:30], flat(opt)[:30])))
                    continue
                if not vi.strip() or KANA.search(vi):
                    problems.append((key, "VN trống hoặc còn kana"))
                    continue
                if ("[主人公]" in opt) != ("[主人公]" in vi):
                    problems.append((key, "[主人公] chỉ có một bên"))
                    continue
                new_opts[i] = vi.replace("\n", " ")
            if new_opts != opts:
                doc["target"] = new_opts
                new_raw = json.dumps(doc, ensure_ascii=False, separators=(",", ":"))
                oj, nj = json.dumps(cur_raw, ensure_ascii=False), json.dumps(new_raw, ensure_ascii=False)
                if out.count(oj) != 1:
                    raise SystemExit("selText sID=%d [%d] khớp %d lần" % (sid, k, out.count(oj)))
                out = out.replace(oj, nj)
                n_sel += sum(1 for a, b in zip(opts, new_opts) if a != b)
        if cur_s != script:
            oj, nj = json.dumps(script, ensure_ascii=False), json.dumps(cur_s, ensure_ascii=False)
            if out.count(oj) != 1:
                raise SystemExit("scriptText target[%d] khớp %d lần" % (ti, out.count(oj)))
            out = out.replace(oj, nj)
        # nameplate
        names = list(t["talkName"])
        new_names = [NAMEPLATE.get(n, n) if n else n for n in names]
        unknown = sorted({n for n in names if n and n not in NAMEPLATE})
        if unknown:
            problems.append(("%d/talkName" % sid, "nameplate chưa có trong NAMEPLATE: %s" % unknown))
        elif new_names != names:
            out = replace_array(out, names, new_names, "talkName target[%d]" % ti)
            n_name += sum(1 for a, b in zip(names, new_names) if a != b)

    print("ghi %d/%d câu, %d phương án lựa chọn, đổi %d nameplate, mirror scriptText %d (không khớp %d)"
          % (n_ok, sum(1 for t in data["target"] for s in t["text"] if s.strip()), n_sel, n_name, mirrored, mirror_failed))
    for k, why in problems:
        print("  !! %-16s %s" % (k, why))
    if problems:
        raise SystemExit("%d ô bị chốt chặn — không ghi" % len(problems))

    after = json.loads(out.lstrip("\ufeff"))
    for t, ta in zip(data["target"], after["target"]):
        assert ta["loadLine"] == t["loadLine"] and ta["scriptText_Line"] == t["scriptText_Line"]
        assert ta["selLine"] == t["selLine"] and len(ta["selText"]) == len(t["selText"])
        assert len(ta["text"]) == len(t["text"])
        assert ta["isDefaultNameAdjust"] == t["isDefaultNameAdjust"]
        for s0, s1 in zip(t["selText"], ta["selText"]):
            assert bool(s0) == bool(s1)
            if s0:
                assert len(json.loads(s0)["target"]) == len(json.loads(s1)["target"])
    print("kiểm tra: loadLine/scriptText_Line/selLine/cờ tên nguyên vẹn, selText giữ đúng số phương án")

    envj, dj, rawj = load_text_asset(os.path.join(STOCK_AOC, JSONB), DLCDATA)
    outj = rawj
    for jp, vn in CHARA_NAME.items():
        a, b = '"jp": "%s"' % jp, '"jp": "%s"' % vn
        assert outj.count(a) == 1, (jp, outj.count(a))
        outj = outj.replace(a, b)
    json.loads(outj.lstrip("\ufeff"))
    print("%s: %d charaName -> La-tinh" % (DLCDATA, len(CHARA_NAME)))

    if not APPLY:
        print("\nCHẠY THỬ — thêm --apply để ghi %s" % WORK_AOC)
        return
    for rel, envx, dx, blob in ((SCEN, env, d, bom + out.lstrip("\ufeff")),
                                (JSONB, envj, dj, outj)):
        dst = os.path.join(WORK_AOC, rel)
        os.makedirs(os.path.dirname(dst), exist_ok=True)
        dx.m_Script = blob
        dx.save()
        with open(dst, "wb") as f:
            f.write(envx.file.save(packer="lz4"))
        print("đã ghi %s (%d B)" % (dst, os.path.getsize(dst)))
    check_work()


if __name__ == "__main__":
    main()
