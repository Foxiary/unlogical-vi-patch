# -*- coding: utf-8 -*-
"""Xuất 80 mục DictionaryData ra một tab .xlsx đúng khuôn `apply_sheet_cells.py` đọc được.

DictionaryData là asset duy nhất còn chữ để dịch mà **sheet không phủ**: snapshot (42)
có 142 tab (INDEX + Bảng Xưng Hô + 131 tab `sd_*` + 9 tab `*Data`), không tab nào cho
từ điển — quét cả workbook bằng chuỗi đặc trưng của mục 102 (`プログラムで目的を達成するた`)
ra 0 ô. Toàn bộ 80 mục được dịch thẳng trên file bằng `fix_dictionary_*.py`, nên chưa
bao giờ có bản trên sheet để đối chiếu hay để người khác sửa.

Khuôn ra giống các tab `*Data` đang có: `ID | Japanese | Vietnamese`, id dạng
`DictionaryData/<nhãn>/id<no>`, nhãn theo lối `news_title`/`news_body` sẵn có:

    DictionaryData/dic_title/id102     title   80 hàng
    DictionaryData/dic_ruby/id102      ruby    49 hàng — 31 mục không có field này
    DictionaryData/dic_body/id102      text    80 hàng, tổng 947 dòng ngắt tay

**Ngắt dòng để nguyên Alt+Enter, KHÔNG đổi thành `\n` văn bản.** Tab `*Data` là bên sở
hữu bố cục: `read_sheet()` đổi Alt+Enter thành dấu `\n` văn bản rồi `expand_breaks()`
mở lại, nên round-trip không mất gì — đã đo là 80 mục không có `　` thụt lề lẫn space
cạnh ngắt dòng, hai thứ duy nhất mà regex gộp khoảng trắng của `read_sheet` sẽ ăn mất.
Thân mục vốn do `fix_dictionary_wrap.py` ngắt tay theo khung, làm phẳng là mất hết.

`category` (あ/か/さ…) ra cột D làm thông tin, **không** phải cột dịch: đó là khoá phân
tab あかさたな của màn ARCHIVE. Cột D nằm ngoài vùng `read_sheet` đọc (A/B/C) nên vô hại.

Ruby rỗng (`""`, 21 mục) vẫn ra hàng: build rỗng + sheet rỗng thì merge coi là "đã có
bản mới" và bỏ qua, nên hàng đó chỉ là chỗ trống để điền sau. Mục KHÔNG có field `ruby`
thì không ra hàng — merge sẽ báo "không có field 'ruby'".

    python tools\export_dictionary_sheet.py [--out=D:\Downloads\DictionaryData_sheet.xlsx]
"""
import io
import json
import os
import sys

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import UnityPy                                          # noqa: E402
from openpyxl import Workbook                           # noqa: E402
from openpyxl.styles import Alignment, Font              # noqa: E402

VN_BUNDLE = os.path.join(ROOT, "romfs", "Data", "StreamingAssets", "json", "json")
JP_BUNDLE = r"D:\Downloads\UNLOGICAL_v2\Data\StreamingAssets\json\json"
OUT = r"D:\Downloads\DictionaryData_sheet.xlsx"

# `title` -> nhãn sheet.  Thứ tự này cũng là thứ tự hàng trong mỗi mục.
FIELDS = [("title", "dic_title"), ("ruby", "dic_ruby"), ("text", "dic_body")]


def arg(name, default=None):
    for a in sys.argv:
        if a.startswith("--%s=" % name):
            return a.split("=", 1)[1]
    return default


def load_dict(path):
    env = UnityPy.load(path)
    for o in env.objects:
        if o.type.name != "TextAsset":
            continue
        d = o.read()
        if d.m_Name != "DictionaryData":
            continue
        raw = d.m_Script
        if not isinstance(raw, str):
            raw = bytes(raw).decode("utf-8")
        return json.loads(raw.lstrip("\ufeff"))
    raise SystemExit("không thấy DictionaryData trong " + path)


def main():
    out_path = arg("out", OUT)
    vn = {i["no"]: i for i in load_dict(VN_BUNDLE)["data"]}
    jp = {i["no"]: i for i in load_dict(JP_BUNDLE)["data"]}

    wb = Workbook()
    ws = wb.active
    ws.title = "DictionaryData"
    ws.append(["ID", "Japanese", "Vietnamese", "category (khoá — không dịch)"])
    for c in ws[1]:
        c.font = Font(bold=True)

    rows, skipped = 0, 0
    for no in sorted(vn):
        it, jt = vn[no], jp.get(no, {})
        for field, label in FIELDS:
            if field not in it:
                skipped += 1
                continue
            ws.append(["DictionaryData/%s/id%d" % (label, no),
                       (jt.get(field) or {}).get("jp", ""),
                       it[field]["jp"],
                       it["category"] if field == "title" else ""])
            rows += 1

    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 44
    ws.column_dimensions["D"].width = 26
    wrap = Alignment(wrap_text=True, vertical="top")
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=3):
        for c in row:
            c.alignment = wrap
    ws.freeze_panes = "A2"

    wb.save(out_path)
    print("đã ghi %s  (%d hàng, bỏ %d hàng ruby vì mục không có field)"
          % (out_path, rows, skipped))

    # Đọc lại từ disk để xác nhận — kể cả ngắt dòng, vì đó là thứ dễ mất nhất.
    from openpyxl import load_workbook
    ws2 = load_workbook(out_path, read_only=True)["DictionaryData"]
    got = {}
    for r in ws2.iter_rows(min_row=2, values_only=True):
        got[r[0]] = r[2]
    bad = 0
    for no in sorted(vn):
        for field, label in FIELDS:
            if field not in vn[no]:
                continue
            k = "DictionaryData/%s/id%d" % (label, no)
            want = vn[no][field]["jp"]
            if (got.get(k) or "") != want:
                bad += 1
                print("!! %s lệch: %r != %r" % (k, got.get(k), want))
    nl = sum((v or "").count("\n") for k, v in got.items() if "/dic_body/" in k)
    print("đọc lại: %d hàng, %d lệch, %d ngắt dòng trong dic_body (asset: %d)"
          % (len(got), bad,
             nl, sum(vn[n]["text"]["jp"].count("\n") for n in vn)))
    if bad:
        raise SystemExit(1)


main()
